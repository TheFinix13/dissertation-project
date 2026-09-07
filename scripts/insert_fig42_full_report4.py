#!/usr/bin/env python3
"""Insert Figure 4.2 (chrono split) in §4.2.3 of Full report - 4.docx."""
from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph

SRC = Path(
    "/Users/the1finix/Documents/University of Surrey/Feb 2026 Courses/"
    "Surrey Courses/Final Dissertation/Dissertation Drafts/Full report - 4.docx"
)
OUT = Path(
    "/Users/the1finix/Documents/University of Surrey/Feb 2026 Courses/"
    "Surrey Courses/Final Dissertation/Dissertation Drafts/"
    "Full report - 4 - with Fig 4.1 and 4.2.docx"
)
FIG = Path(__file__).resolve().parents[1] / "latex/final_dissertation/figs/chrono_split.png"

ANCHOR = "The split is:"
BRIDGE = "Figure 4.2 shows the split."
CAPTION = (
    "Figure 4.2: Chronological three-way split of the 96 monthly episodes: "
    "60 training months, 12 validation months and 24 held-out test months."
)


def set_run_font(run, *, bold=False, italic=False, size=11):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def add_run(p, text, **kw):
    r = p.add_run(text)
    set_run_font(r, **kw)
    return r


def insert_paragraph_after(paragraph: Paragraph, text: str = "") -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if text:
        add_run(new_para, text, size=11)
    return new_para


def figure_already_present(doc: Document) -> bool:
    return any(p.text.strip().startswith("Figure 4.2:") for p in doc.paragraphs)


def main() -> None:
    if not FIG.exists():
        raise FileNotFoundError(f"figure PNG missing: {FIG}")

    doc = Document(str(SRC))
    if figure_already_present(doc):
        print("Figure 4.2 already present; writing copy anyway.")

    anchor_idx = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip() == ANCHOR)
    prev = doc.paragraphs[anchor_idx]

    bridge = insert_paragraph_after(prev, BRIDGE)
    bridge.paragraph_format.space_after = Pt(6)

    img = insert_paragraph_after(bridge, "")
    img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    img.paragraph_format.space_before = Pt(6)
    img.paragraph_format.space_after = Pt(6)
    img.add_run().add_picture(str(FIG), width=Cm(15.5))

    cap = insert_paragraph_after(img, "")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(cap, CAPTION, italic=True, size=10)
    cap.paragraph_format.space_after = Pt(12)

    doc.save(str(OUT))
    try:
        shutil.copy2(OUT, SRC)
        print(f"updated {SRC}")
    except Exception as exc:
        print(f"could not overwrite {SRC} ({exc}); saved to {OUT}")


if __name__ == "__main__":
    main()
