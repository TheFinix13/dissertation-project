#!/usr/bin/env python3
"""Insert Algorithms 1–3 (REINFORCE, DQN, MaskablePPO) under §3.10 in Full Report - 4.docx."""
from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, Twips
from docx.text.paragraph import Paragraph

SRC = Path(
    "/Users/the1finix/Documents/University of Surrey/Feb 2026 Courses/"
    "Surrey Courses/Final Dissertation/Dissertation Drafts/Full Report - 4.docx"
)
OUT = Path(
    "/Users/the1finix/Documents/University of Surrey/Feb 2026 Courses/"
    "Surrey Courses/Final Dissertation/Dissertation Drafts/"
    "Full Report - 4 - with algorithms.docx"
)


def set_run_font(run, *, bold=False, italic=False, size=10):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def add_run(p, text, **kw):
    r = p.add_run(text)
    set_run_font(r, **kw)
    return r


def insert_paragraph_after(paragraph: Paragraph, text: str = "") -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if text:
        add_run(new_para, text, size=11)
    return new_para


def set_para_spacing(p, before=0, after=0, left=0):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if left:
        pf.left_indent = Cm(left)


def set_cell_borders(cell, *, top=None, bottom=None, left=None, right=None):
    """Border spec: (sz, val, color) or None for nil."""
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag == qn("w:tcBorders"):
            tcPr.remove(child)
    borders = OxmlElement("w:tcBorders")
    for edge, spec in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{edge}")
        if spec is None:
            el.set(qn("w:val"), "nil")
        else:
            sz, val, color = spec
            el.set(qn("w:val"), val)
            el.set(qn("w:sz"), str(sz))
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), color)
        borders.append(el)
    tcPr.append(borders)


def clear_cell(cell):
    tc = cell._tc
    for child in list(tc):
        if child.tag == qn("w:p"):
            tc.remove(child)
    # leave no paragraphs; caller adds them


def add_cell_paragraph(cell, *, before=0, after=0, left_cm=0.0, align=None):
    p = OxmlElement("w:p")
    cell._tc.append(p)
    para = Paragraph(p, cell)
    set_para_spacing(para, before=before, after=after, left=left_cm)
    if align is not None:
        para.alignment = align
    return para


def algorithm_line(cell, number: int | None, parts, *, indent=0):
    """parts: list of (text, bold, italic) tuples. number=None for blank/end lines."""
    left = 0.35 + 0.45 * indent
    p = add_cell_paragraph(cell, before=0, after=1, left_cm=left)
    if number is not None:
        add_run(p, f"{number}:  ", bold=False, size=10)
    for text, bold, italic in parts:
        add_run(p, text, bold=bold, italic=italic, size=10)
    return p


KW = lambda s: (s, True, False)   # bold keyword
TX = lambda s: (s, False, False)
IT = lambda s: (s, False, True)


def build_algorithm_table(doc, anchor_para: Paragraph, caption: str, lines):
    """Insert a one-cell booktabs-style algorithm box after anchor_para.

    lines: list of (number|None, indent, parts)
    Returns the paragraph after the inserted table (for chaining).
    """
    # spacer
    spacer = insert_paragraph_after(anchor_para, "")
    set_para_spacing(spacer, before=6, after=0)

    table = doc.add_table(rows=1, cols=1)
    table.autofit = True
    # move table to sit after spacer
    spacer._p.addnext(table._tbl)

    cell = table.rows[0].cells[0]
    clear_cell(cell)
    # outer booktabs: thick top + bottom
    set_cell_borders(
        cell,
        top=(18, "single", "000000"),
        bottom=(18, "single", "000000"),
        left=None,
        right=None,
    )

    # caption
    cap = add_cell_paragraph(cell, before=4, after=4, align=WD_ALIGN_PARAGRAPH.LEFT)
    add_run(cap, caption, bold=True, size=10)

    # rule under caption: thin bottom border on a one-line empty para via shading trick —
    # emulate with a paragraph of underscores? Better: nested single-row table for rule.
    rule = add_cell_paragraph(cell, before=0, after=4)
    # draw a horizontal line via bottom border on this paragraph's pBdr
    pPr = rule._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    pBdr.append(bottom)
    pPr.append(pBdr)

    for number, indent, parts in lines:
        algorithm_line(cell, number, parts, indent=indent)

    # paragraph after the table, used as the next insertion anchor
    new_p = OxmlElement("w:p")
    table._tbl.addnext(new_p)
    after_para = Paragraph(new_p, spacer._parent)
    set_para_spacing(after_para, before=6, after=6)
    return after_para


REINFORCE = [
    (1, 0, [KW("initialise"), TX(" policy network "), IT("π"), TX("θ")]),
    (2, 0, [KW("for"), TX(" iteration = 1, 2, … until step budget exhausted "), KW("do")]),
    (3, 1, [TX("sample a training month; reset environment ("), IT("C"), TX("₀ cash, no position)")]),
    (4, 1, [KW("for"), TX(" "), IT("t"), TX(" = 0, …, "), IT("T"), TX(" − 1 "), KW("do")]),
    (5, 2, [TX("compute legal-action mask; set illegal logits to −∞")]),
    (6, 2, [TX("sample "), IT("a"), TX("ₜ ∼ "), IT("π"), TX("θ(· ∣ "), IT("s"), TX("ₜ); execute; store (log "), IT("π"), TX("θ("), IT("a"), TX("ₜ∣"), IT("s"), TX("ₜ), "), IT("r"), TX("ₜ)")]),
    (7, 1, [KW("end for")]),
    (8, 1, [TX("compute rewards-to-go "), IT("G"), TX("ₜ; standardise within the episode")]),
    (9, 1, [IT("L"), TX("(θ) ← −∑ₜ log "), IT("π"), TX("θ("), IT("a"), TX("ₜ∣"), IT("s"), TX("ₜ) "), IT("G"), TX("ₜ")]),
    (10, 1, [TX("one Adam step on ∇θ"), IT("L")]),
    (11, 0, [KW("end for")]),
]

DQN = [
    (1, 0, [KW("initialise"), TX(" "), IT("Q"), TX("ϕ; target "), IT("Q"), TX("ϕ⁻ ← "), IT("Q"), TX("ϕ; replay buffer "), IT("D")]),
    (2, 0, [KW("for"), TX(" environment step = 1, 2, … until step budget exhausted "), KW("do")]),
    (3, 1, [TX("with prob. ε: random legal action; else "), IT("a"), TX("ₜ = arg max"), IT("a"), TX(" legal "), IT("Q"), TX("ϕ("), IT("s"), TX("ₜ, "), IT("a"), TX(")")]),
    (4, 1, [TX("execute "), IT("a"), TX("ₜ; store ("), IT("s"), TX("ₜ, "), IT("a"), TX("ₜ, "), IT("r"), TX("ₜ, "), IT("s"), TX("ₜ₊₁, maskₜ₊₁, "), IT("d"), TX("ₜ) in "), IT("D")]),
    (5, 1, [TX("sample minibatch; form masked targets "), IT("y"), TX(" per (3.16); Adam step on "), IT("L"), TX("(ϕ)")]),
    (6, 1, [TX("every "), IT("K"), TX(" steps: ϕ⁻ ← ϕ")]),
    (7, 0, [KW("end for")]),
]

# Matches experiments/final_v2/harness.py + MonthSampler + sb3_contrib.MaskablePPO
# defaults used in this project: n_steps=2048, batch=64, n_epochs=10, clip ε=0.2, GAE.
PPO = [
    (1, 0, [KW("initialise"), TX(" actor "), IT("π"), TX("θ and critic "), IT("V"), TX("ψ (MaskablePPO, MLP 128–128)")]),
    (2, 0, [KW("for"), TX(" iteration = 1, 2, … until step budget exhausted "), KW("do")]),
    (3, 1, [KW("for"), TX(" step = 1, …, "), IT("N"), TX(" "), KW("do")]),
    (4, 2, [KW("if"), TX(" episode terminated: sample a training month; reset ("), IT("C"), TX("₀ cash, no position)")]),
    (5, 2, [TX("observe "), IT("s"), TX("ₜ and legal-action mask; sample "), IT("a"), TX("ₜ ∼ "), IT("π"), TX("θ(· ∣ "), IT("s"), TX("ₜ) over legal actions")]),
    (6, 2, [TX("execute; store ("), IT("s"), TX("ₜ, "), IT("a"), TX("ₜ, "), IT("r"), TX("ₜ, "), IT("V"), TX("ψ("), IT("s"), TX("ₜ), log "), IT("π"), TX("θ("), IT("a"), TX("ₜ∣"), IT("s"), TX("ₜ))")]),
    (7, 1, [KW("end for")]),
    (8, 1, [TX("estimate advantages "), IT("Â"), TX("ₜ by GAE; form returns")]),
    (9, 1, [KW("for"), TX(" epoch = 1, …, "), IT("K"), TX(" "), KW("do")]),
    (10, 2, [KW("for"), TX(" each minibatch from the rollout "), KW("do")]),
    (11, 3, [IT("ρ"), TX("ₜ ← "), IT("π"), TX("θ("), IT("a"), TX("ₜ∣"), IT("s"), TX("ₜ) / "), IT("π"), TX("θ_old("), IT("a"), TX("ₜ∣"), IT("s"), TX("ₜ)")]),
    (12, 3, [TX("Adam step on "), IT("L"), TX("CLIP(θ) + value loss")]),
    (13, 2, [KW("end for")]),
    (14, 1, [KW("end for")]),
    (15, 0, [KW("end for")]),
]


def main():
    doc = Document(str(SRC))

    anchor = None
    for p in doc.paragraphs:
        if p.text.strip() == "3.10 Algorithms":
            anchor = p
            break
    if anchor is None:
        raise SystemExit("Could not find heading '3.10 Algorithms'")

    # Short intro under the heading (Talha-style: what this section is)
    intro = insert_paragraph_after(anchor, "")
    set_para_spacing(intro, before=6, after=6)
    add_run(
        intro,
        "The three learning procedures used in this dissertation are summarised "
        "below as implemented. REINFORCE and Deep Q-learning are the from-scratch "
        "methods; MaskablePPO is the library comparator that applies the same "
        "legal-action mask.",
        size=11,
    )

    after = build_algorithm_table(
        doc, intro,
        "Algorithm 1: REINFORCE for the trading MDP (as implemented)",
        REINFORCE,
    )
    after = build_algorithm_table(
        doc, after,
        "Algorithm 2: Deep Q-learning for the trading MDP (as implemented)",
        DQN,
    )
    after = build_algorithm_table(
        doc, after,
        "Algorithm 3: MaskablePPO for the trading MDP (as implemented)",
        PPO,
    )

    note = after
    add_run(
        note,
        "Notes. N = 2048 is the Stable-Baselines3 rollout length used here; "
        "K = 10 is the number of epochs over each rollout. Month sampling and "
        "action masking use the same environment interface as REINFORCE and DQN, "
        "so all three algorithms face an identical action space at every step.",
        italic=True,
        size=10,
    )
    set_para_spacing(note, before=2, after=8)

    bridge = insert_paragraph_after(note, "")
    add_run(
        bridge,
        "The next section states why all three methods are retained in the "
        "experimental comparison.",
        size=11,
    )
    set_para_spacing(bridge, before=6, after=6)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(f"Wrote {OUT}")

    # Try updating the original if unlocked
    try:
        shutil.copy2(OUT, SRC)
        print(f"Also updated {SRC}")
    except Exception as e:
        print(f"Could not overwrite original (likely open in Word): {e}")
        print("Close Full Report - 4.docx without saving, then open the new file.")


if __name__ == "__main__":
    main()
