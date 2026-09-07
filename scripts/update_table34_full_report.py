#!/usr/bin/env python3
"""Update Table 3.4 and ATR/MA prose in Full Report .docx (implementation-accurate)."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm

SRC = Path(
    "/Users/the1finix/Documents/University of Surrey/Feb 2026 Courses/"
    "Surrey Courses/Final Dissertation/Dissertation Drafts/Full Report .docx"
)
OUT = Path(
    "/Users/the1finix/Documents/University of Surrey/Feb 2026 Courses/"
    "Surrey Courses/Final Dissertation/Dissertation Drafts/"
    "Full Report - Table 3.4 verified.docx"
)


def set_run_font(run, bold=False, italic=False, size=10, sub=False, sup=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.subscript = sub
    run.font.superscript = sup


def add_text(p, text, **kw):
    r = p.add_run(text)
    set_run_font(r, **kw)
    return r


def clear_cell(cell):
    for p in cell.paragraphs:
        for r in list(p.runs):
            r._element.getparent().remove(r._element)
    while len(cell.paragraphs) > 1:
        cell.paragraphs[-1]._element.getparent().remove(cell.paragraphs[-1]._element)


def cell_para(cell):
    clear_cell(cell)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    return p


def set_cell_borders(cell, top=None, bottom=None):
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag == qn("w:tcBorders"):
            tcPr.remove(child)
    borders = OxmlElement("w:tcBorders")
    for edge, spec in (("top", top), ("left", None), ("bottom", bottom), ("right", None)):
        el = OxmlElement(f"w:{edge}")
        if spec is None:
            el.set(qn("w:val"), "nil")
        else:
            sz, val, color = spec
            el.set(qn("w:val"), val)
            el.set(qn("w:sz"), str(sz))
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), color)
        borders.append(el)
    tcPr.append(borders)


def apply_booktabs(table):
    n = len(table.rows)
    for ri, row in enumerate(table.rows):
        for cell in row.cells:
            top = bottom = None
            if ri == 0:
                top = (18, "single", "000000")
                bottom = (12, "single", "000000")
            elif ri == n - 1:
                bottom = (18, "single", "000000")
            set_cell_borders(cell, top=top, bottom=bottom)


def set_col_widths(table, widths_cm):
    for row in table.rows:
        for cell, w in zip(row.cells, widths_cm):
            cell.width = Cm(w)


def find_table_34(doc: Document):
    for t in doc.tables:
        hdr = [c.text.strip() for c in t.rows[0].cells]
        if hdr == ["#", "Feature", "Definition", "What it tells the agent"]:
            return t
    raise RuntimeError("Table 3.4 not found")


def rebuild_table_34(table):
    while len(table.rows) < 14:
        table.add_row()
    while len(table.rows) > 14:
        table._tbl.remove(table.rows[-1]._tr)

    set_col_widths(table, [0.9, 2.0, 5.2, 6.5])

    for cell, label in zip(
        table.rows[0].cells,
        ["#", "Feature", "Definition", "What it tells the agent"],
    ):
        add_text(cell_para(cell), label, bold=True, size=10)

    def section_row(row, title):
        row.cells[0].merge(row.cells[3])
        add_text(cell_para(row.cells[0]), title, italic=True, bold=True, size=10)

    def feature_row(row, num, feat_fn, def_fn, tells):
        add_text(cell_para(row.cells[0]), str(num), size=10)
        feat_fn(cell_para(row.cells[1]))
        def_fn(cell_para(row.cells[2]))
        add_text(cell_para(row.cells[3]), tells, size=10)

    section_row(table.rows[1], "Market")

    feature_row(
        table.rows[2], 1,
        lambda p: (add_text(p, "ΔP", italic=True), add_text(p, "t", italic=True, sub=True)),
        lambda p: (
            add_text(p, "(", size=10),
            add_text(p, "P", italic=True), add_text(p, "t", italic=True, sub=True),
            add_text(p, " − ", size=10),
            add_text(p, "P", italic=True), add_text(p, "t−1", italic=True, sub=True),
            add_text(p, ") / ", size=10),
            add_text(p, "P", italic=True), add_text(p, "t−1", italic=True, sub=True),
        ),
        "the immediate one-step price return",
    )

    feature_row(
        table.rows[3], 2,
        lambda p: (add_text(p, "mom", italic=True), add_text(p, "t", italic=True, sub=True)),
        lambda p: (
            add_text(p, "P", italic=True), add_text(p, "t", italic=True, sub=True),
            add_text(p, " / ", size=10),
            add_text(p, "P", italic=True), add_text(p, "t−k", italic=True, sub=True),
            add_text(p, " − 1  (k = 5)", size=10),
        ),
        "whether recent movement is part of a short-term trend",
    )

    feature_row(
        table.rows[4], 3,
        lambda p: (add_text(p, "vol", italic=True), add_text(p, "t", italic=True, sub=True)),
        lambda p: (
            add_text(p, "population std. of ", size=10),
            add_text(p, "r", italic=True), add_text(p, "t−i", italic=True, sub=True),
            add_text(p, " for i = 0,…,k−1; k = 5; ddof = 0", size=10),
        ),
        "how variable the market has been over the last k days",
    )

    feature_row(
        table.rows[5], 4,
        lambda p: (add_text(p, "gap", italic=True), add_text(p, "t", italic=True, sub=True)),
        lambda p: (
            add_text(p, "P", italic=True), add_text(p, "t", italic=True, sub=True),
            add_text(p, " / MA", size=10),
            add_text(p, "m", italic=True, sub=True),
            add_text(p, "(P)", size=10),
            add_text(p, "t", italic=True, sub=True),
            add_text(p, " − 1  (m = 20)", size=10),
        ),
        "how far price sits above or below its recent average",
    )

    feature_row(
        table.rows[6], 5,
        lambda p: (add_text(p, "rng", italic=True), add_text(p, "t", italic=True, sub=True)),
        lambda p: (
            add_text(p, "ATR", size=10),
            add_text(p, "n,t", italic=True, sub=True),
            add_text(p, " / ", size=10),
            add_text(p, "P", italic=True), add_text(p, "t", italic=True, sub=True),
            add_text(p, "  (n = 14)", size=10),
        ),
        "typical daily range, including gaps; not visible from closes alone",
    )

    section_row(table.rows[7], "Portfolio")

    feature_row(
        table.rows[8], 6,
        lambda p: (add_text(p, "c", italic=True), add_text(p, "t", italic=True, sub=True)),
        lambda p: (
            add_text(p, "C", italic=True), add_text(p, "t", italic=True, sub=True),
            add_text(p, " / ", size=10),
            add_text(p, "C", italic=True), add_text(p, "0", italic=True, sub=True),
        ),
        "how much cash is available to deploy",
    )

    feature_row(
        table.rows[9], 7,
        lambda p: (add_text(p, "v", italic=True), add_text(p, "t", italic=True, sub=True)),
        lambda p: (
            add_text(p, "h", italic=True), add_text(p, "t", italic=True, sub=True),
            add_text(p, " ", size=10),
            add_text(p, "P", italic=True), add_text(p, "t", italic=True, sub=True),
            add_text(p, " / ", size=10),
            add_text(p, "C", italic=True), add_text(p, "0", italic=True, sub=True),
        ),
        "what the open position is worth",
    )

    section_row(table.rows[10], "Position")

    feature_row(
        table.rows[11], 8,
        lambda p: (add_text(p, "PnL", italic=True), add_text(p, "t", italic=True, sub=True)),
        lambda p: (
            add_text(p, "(", size=10),
            add_text(p, "P", italic=True), add_text(p, "t", italic=True, sub=True),
            add_text(p, " − ", size=10),
            add_text(p, "P̄", italic=True), add_text(p, "t", italic=True, sub=True),
            add_text(p, "entry", italic=True, size=8, sub=True),
            add_text(p, ") / ", size=10),
            add_text(p, "P̄", italic=True), add_text(p, "t", italic=True, sub=True),
            add_text(p, "entry", italic=True, size=8, sub=True),
        ),
        "whether the open position is in profit or loss",
    )

    section_row(table.rows[12], "Time")

    feature_row(
        table.rows[13], 9,
        lambda p: (add_text(p, "τ", italic=True), add_text(p, "t", italic=True, sub=True)),
        lambda p: (
            add_text(p, "t", italic=True),
            add_text(p, " / ", size=10),
            add_text(p, "T", italic=True),
            add_text(p, "max", italic=True, sub=True),
            add_text(p, "  ∈ [0, 1]", size=10),
        ),
        "how far through the trading episode has elapsed",
    )

    apply_booktabs(table)


def rewrite_paragraph(p, lines: list[tuple[str, dict]]):
    """Replace paragraph content with formatted runs."""
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    for text, kw in lines:
        if text:
            add_text(p, text, **kw)


def update_prose(doc: Document):
    replacements = {
        "Where ": [
            ("Where mom", {"italic": True}),
            ("t", {"italic": True, "sub": True}),
            (" = P", {"size": 10}),
            ("t", {"italic": True, "sub": True}),
            (" / P", {"size": 10}),
            ("t−k", {"italic": True, "sub": True}),
            (" − 1, with k = 5.", {"size": 10}),
        ],
        "Where  represents the standard deviation": [
            ("Where σ", {"italic": True}),
            ("k", {"italic": True, "sub": True}),
            ("(r)", {"italic": True}),
            ("t", {"italic": True, "sub": True}),
            (" is the population standard deviation (ddof = 0) of the one-step returns "
             "r", {"size": 10}),
            ("t−i", {"italic": True, "sub": True}),
            (" = (P", {"size": 10}),
            ("t−i", {"italic": True, "sub": True}),
            (" − P", {"size": 10}),
            ("t−i−1", {"italic": True, "sub": True}),
            (") / P", {"size": 10}),
            ("t−i−1", {"italic": True, "sub": True}),
            (" for i = 0,…,k−1, with k = 5.", {"size": 10}),
        ],
        "where  represents the moving average": [
            ("where MA", {"size": 10}),
            ("m,t", {"italic": True, "sub": True}),
            ("(P) is the simple moving average of the closing price over the previous "
             "m timesteps:", {"size": 10}),
        ],
        "Using the standard Wilder smoothing procedure": [
            ("In this implementation, ATR", {"size": 10}),
            ("n,t", {"italic": True, "sub": True}),
            (" is the arithmetic mean of the last n True Range values (a rolling "
             "average, not Wilder's recursive smoother):", {"size": 10}),
        ],
    }

    # insert MA formula paragraph after moving-average intro if missing
    ma_formula_idx = None
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if t.startswith("where MA"):
            ma_formula_idx = i
        for key, lines in replacements.items():
            if t.startswith(key) or key in t:
                rewrite_paragraph(p, lines)
        if "m=10" in t:
            rewrite_paragraph(
                p,
                [(t.replace("m=10", "m=20").replace("the mark-to-market wealth. The", "The"), {"size": 10})],
            )
        if "1 - t / T" in t.replace(" ", ""):
            rewrite_paragraph(
                p,
                [("t / T", {"italic": True}),
                 ("max", {"italic": True, "sub": True}),
                 ("  (episode progress clock; 0 at the first day, approaching 1 near the end)", {"size": 10})],
            )

    if ma_formula_idx is not None:
        nxt = doc.paragraphs[ma_formula_idx + 1].text.strip()
        if not nxt.startswith("MA"):
            p = doc.paragraphs[ma_formula_idx]._p
            new_p = OxmlElement("w:p")
            p.addnext(new_p)
            from docx.text.paragraph import Paragraph
            para = Paragraph(new_p, doc.paragraphs[ma_formula_idx]._parent)
            rewrite_paragraph(
                para,
                [
                    ("MA", {"size": 10}),
                    ("m,t", {"italic": True, "sub": True}),
                    ("(P) = (1/m) Σ", {"size": 10}),
                    ("i=0", {"italic": True, "sub": True}),
                    ("m−1", {"italic": True, "sub": True}),
                    (" P", {"size": 10}),
                    ("t−i", {"italic": True, "sub": True}),
                    ("  (m = 20).", {"size": 10}),
                ],
            )

    # ATR rolling mean formula after Wilder replacement paragraph
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip().startswith("In this implementation, ATR"):
            nxt = doc.paragraphs[i + 1].text.strip() if i + 1 < len(doc.paragraphs) else ""
            if not nxt.startswith("ATR"):
                new_p = OxmlElement("w:p")
                p._p.addnext(new_p)
                from docx.text.paragraph import Paragraph
                para = Paragraph(new_p, p._parent)
                rewrite_paragraph(
                    para,
                    [
                        ("ATR", {"size": 10}),
                        ("n,t", {"italic": True, "sub": True}),
                        (" = (1/n) Σ", {"size": 10}),
                        ("i=0", {"italic": True, "sub": True}),
                        ("n−1", {"italic": True, "sub": True}),
                        (" TR", {"size": 10}),
                        ("t−i", {"italic": True, "sub": True}),
                        ("  (n = 14).", {"size": 10}),
                    ],
                )
            break


def main():
    doc = Document(str(SRC))
    rebuild_table_34(find_table_34(doc))
    update_prose(doc)
    doc.save(str(OUT))
    print(f"Saved: {OUT}")
    try:
        doc.save(str(SRC))
        print(f"Also updated: {SRC}")
    except PermissionError:
        print("Original locked — close Word and re-run to update Full Report .docx")


if __name__ == "__main__":
    main()
